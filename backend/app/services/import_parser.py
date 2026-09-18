from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from codecs import getincrementaldecoder
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from ..core.config import settings


class ImportValidationError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(slots=True)
class RowError:
    row_no: int | None
    code: str
    message: str
    raw_data: dict[str, Any] | None = None


@dataclass(slots=True)
class ParsedImport:
    data_type: str
    source_file: str
    business_date: date | None
    rows: list[dict[str, Any]] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    warnings: list[RowError] = field(default_factory=list)
    header_mapping: dict[str, str] = field(default_factory=dict)
    source_row_count: int = 0
    skipped_rows: int = 0

    @property
    def valid_rows(self) -> int:
        return len(self.rows)

    @property
    def error_rows(self) -> int:
        return len(self.errors)

    @property
    def warning_rows(self) -> int:
        return len(self.warnings)

    def summary(self) -> dict[str, Any]:
        return {
            "data_type": self.data_type,
            "source_file": self.source_file,
            "business_date": self.business_date.isoformat() if self.business_date else None,
            "source_rows": self.source_row_count,
            "valid_rows": self.valid_rows,
            "error_rows": self.error_rows,
            "warning_rows": self.warning_rows,
            "skipped_rows": self.skipped_rows,
            "header_mapping": self.header_mapping,
        }


SCHEMAS: dict[str, dict[str, Any]] = {
    "product": {
        "required": ["merchant_code", "product_name"],
        "aliases": {
            "merchant_code": ["商家编码", "商品编码", "货品编码"],
            "product_name": ["货品名称", "商品名称", "商品名"],
            "spec": ["规格", "规格名称"],
            "brand": ["品牌"],
            "category": ["分类", "商品分类"],
            "barcode": ["条码", "商品条码"],
        },
    },
    "sales": {
        "required": ["shop_name", "warehouse_name", "merchant_code", "product_name", "sales_qty", "avg_price"],
        "aliases": {
            "shop_name": ["店铺", "店铺名称"],
            "warehouse_name": ["仓库", "仓库名称"],
            "merchant_code": ["商家编码", "商品编码", "货品编码"],
            "product_name": ["商品名称", "货品名称", "商品名"],
            "sales_qty": ["销量", "实际销售量", "销售量", "销售数量"],
            "avg_price": ["均价", "平均售价", "平均价格"],
        },
    },
    "inventory": {
        "required": ["warehouse_name", "merchant_code", "product_name", "stock_qty", "production_date", "expire_date"],
        "aliases": {
            "warehouse_name": ["仓库", "仓库名称"],
            "merchant_code": ["商家编码", "商品编码", "货品编码"],
            "product_name": ["商品名称", "货品名称", "商品名"],
            "stock_qty": ["库存数量", "正常库存", "数量", "库存"],
            "production_date": ["生产日期", "生产时间"],
            "expire_date": ["过期日期", "有效期至", "失效日期"],
            "source_snapshot_date": ["库存日期", "快照日期", "统计日期"],
        },
    },
    "aging": {
        "required": ["warehouse_name", "merchant_code", "product_name", "stock_qty", "aging_days"],
        "aliases": {
            "warehouse_name": ["仓库", "仓库名称"],
            "merchant_code": ["商家编码", "商品编码", "货品编码"],
            "product_name": ["商品名称", "货品名称", "商品名"],
            "stock_qty": ["库存数量", "数量", "库存"],
            "aging_days": ["库龄天数", "库龄[天]", "库龄（天）", "库龄"],
        },
    },
}


_ILLEGAL_XML_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).strip()).lower()


def _header_map(headers: list[Any], data_type: str) -> dict[str, int]:
    schema = SCHEMAS[data_type]
    normalized = {_norm_header(v): i for i, v in enumerate(headers) if _norm_header(v)}
    mapping: dict[str, int] = {}
    for canonical, aliases in schema["aliases"].items():
        for alias in aliases:
            key = _norm_header(alias)
            if key in normalized:
                mapping[canonical] = normalized[key]
                break
    missing = [k for k in schema["required"] if k not in mapping]
    if missing:
        display = ", ".join(missing)
        raise ImportValidationError("MISSING_HEADERS", f"缺少必填字段：{display}")
    return mapping


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _ILLEGAL_XML_CHARS_RE.sub("", str(value)).strip()


def _to_decimal(value: Any, field_name: str) -> Decimal:
    if value is None or _clean_text(value) == "":
        raise ImportValidationError("INVALID_NUMBER", f"{field_name} 不能为空")
    try:
        dec = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        raise ImportValidationError("INVALID_NUMBER", f"{field_name} 不是有效数字：{value}")
    # MySQL columns are DECIMAL(18, 4/6); reject non-finite or clearly
    # unrepresentable values before they reach the database transaction.
    if not dec.is_finite() or abs(dec) >= Decimal("1E+12"):
        raise ImportValidationError("INVALID_NUMBER", f"{field_name} 超出允许的数字范围：{value}")
    return dec


def _to_int(value: Any, field_name: str) -> int:
    dec = _to_decimal(value, field_name)
    if dec != dec.to_integral_value():
        raise ImportValidationError("INVALID_INTEGER", f"{field_name} 必须是整数：{value}")
    return int(dec)


def _to_date(value: Any, field_name: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean_text(value)
    if not text:
        raise ImportValidationError("INVALID_DATE", f"{field_name} 不能为空")
    text = text.split(" ")[0]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ImportValidationError("INVALID_DATE", f"{field_name} 日期格式无法识别：{value}")


def _iter_xlsx(path: Path) -> tuple[list[Any], Iterable[tuple[int, list[Any]]]]:
    try:
        with zipfile.ZipFile(path) as archive:
            uncompressed_bytes = sum(info.file_size for info in archive.infolist())
    except (OSError, zipfile.BadZipFile) as exc:
        raise ImportValidationError("INVALID_XLSX", "Excel 文件格式损坏或无法读取") from exc
    if uncompressed_bytes > settings.max_xlsx_uncompressed_bytes:
        limit_mb = settings.max_xlsx_uncompressed_bytes // (1024 * 1024)
        raise ImportValidationError("XLSX_EXPANDED_TOO_LARGE", f"Excel 解压后数据超过 {limit_mb}MB 上限")

    wb = load_workbook(path, read_only=True, data_only=True)
    if len(wb.worksheets) > settings.max_import_worksheets:
        wb.close()
        raise ImportValidationError(
            "TOO_MANY_WORKSHEETS",
            f"Excel 工作表数量超过 {settings.max_import_worksheets} 个上限",
        )
    ws = wb.active
    # 部分旺店通导出的 xlsx 写错了 worksheet dimension（例如 A1:A9035），
    # 但实际有多列数据。read_only 模式下需要重置维度后再扫描。
    if ws.max_column == 1:
        ws.reset_dimensions()
    rows = ws.iter_rows(values_only=True)
    try:
        headers = list(next(rows))
    except StopIteration:
        wb.close()
        raise ImportValidationError("EMPTY_FILE", "Excel 文件为空")

    def iterator() -> Iterable[tuple[int, list[Any]]]:
        try:
            for row_no, row in enumerate(rows, start=2):
                yield row_no, list(row)
        finally:
            wb.close()

    return headers, iterator()


def _csv_encoding(path: Path) -> str:
    with path.open("rb") as source:
        if source.read(3) == b"\xef\xbb\xbf":
            return "utf-8-sig"
    # Validate incrementally so encoding detection does not duplicate the
    # entire upload in memory. UTF-8 is preferred; GB18030 stays compatible.
    for encoding in ("utf-8", "gb18030"):
        decoder = getincrementaldecoder(encoding)(errors="strict")
        try:
            with path.open("rb") as source:
                for chunk in iter(lambda: source.read(1024 * 1024), b""):
                    decoder.decode(chunk)
                decoder.decode(b"", final=True)
            return encoding
        except UnicodeDecodeError:
            continue
    raise ImportValidationError("CSV_ENCODING", "CSV 编码无法识别，请保存为 UTF-8 CSV")


def _iter_csv(path: Path) -> tuple[list[Any], Iterable[tuple[int, list[Any]]]]:
    source = path.open("r", encoding=_csv_encoding(path), newline="")
    reader = csv.reader(source)
    try:
        headers = next(reader)
    except StopIteration:
        source.close()
        raise ImportValidationError("EMPTY_FILE", "CSV 文件为空")
    except Exception:
        source.close()
        raise

    def iterator() -> Iterable[tuple[int, list[Any]]]:
        try:
            for row_no, row in enumerate(reader, start=2):
                yield row_no, row
        finally:
            source.close()

    return list(headers), iterator()


def _source_rows(path: Path) -> tuple[list[Any], Iterable[tuple[int, list[Any]]]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _iter_xlsx(path)
    if suffix == ".csv":
        return _iter_csv(path)
    if suffix == ".xls":
        raise ImportValidationError("XLS_UNSUPPORTED", "旧版 .xls 暂不直接解析，请另存为 .xlsx 或 CSV")
    raise ImportValidationError("UNSUPPORTED_FILE", "只支持 .xlsx / .csv")


def _looks_like_total(row: list[Any]) -> bool:
    for value in row[:4]:
        text = _clean_text(value)
        if text.startswith("合计") or text.startswith("总计"):
            return True
    return False


def _looks_like_aggregate_sales(path: Path) -> bool:
    name = path.name.lower()
    flags = ("30天", "30day", "30-day", "汇总销量", "月销量", "月度销量")
    return any(flag in name for flag in flags)


def parse_import_file(
    path: str | Path,
    data_type: str,
    business_date: date | None = None,
    *,
    reject_aggregate_sales: bool = True,
) -> ParsedImport:
    if data_type not in SCHEMAS:
        raise ImportValidationError("INVALID_DATA_TYPE", f"不支持的数据类型：{data_type}")

    file_path = Path(path)
    if not file_path.exists():
        raise ImportValidationError("FILE_NOT_FOUND", f"文件不存在：{file_path}")

    if data_type in {"sales", "inventory", "aging"} and business_date is None:
        raise ImportValidationError("BUSINESS_DATE_REQUIRED", f"{data_type} 导入必须由导入者选择业务日期")

    if data_type == "sales" and reject_aggregate_sales and _looks_like_aggregate_sales(file_path):
        raise ImportValidationError(
            "AGGREGATE_SALES_NOT_DAILY",
            "检测到 30 天/汇总销量文件。sales_daily 只能导入逐日销量；请使用逐日文件目录导入，不能把 30 天汇总冒充某一天。",
        )

    headers, rows_iter = _source_rows(file_path)
    try:
        mapping = _header_map(headers, data_type)
    except Exception:
        close = getattr(rows_iter, "close", None)
        if close:
            close()
        raise
    parsed = ParsedImport(data_type=data_type, source_file=file_path.name, business_date=business_date)
    parsed.header_mapping = {key: _clean_text(headers[idx]) for key, idx in mapping.items()}
    seen_names: dict[str, str] = {}

    for row_no, row in rows_iter:
        if not any(_clean_text(v) for v in row):
            parsed.skipped_rows += 1
            continue
        if _looks_like_total(row):
            parsed.skipped_rows += 1
            continue
        parsed.source_row_count += 1
        if parsed.source_row_count > settings.max_import_rows:
            close = getattr(rows_iter, "close", None)
            if close:
                close()
            raise ImportValidationError(
                "TOO_MANY_ROWS",
                f"导入数据超过 {settings.max_import_rows} 行上限，请拆分文件后重新导入",
            )
        raw = {_clean_text(headers[i]) or f"COL_{i+1}": row[i] if i < len(row) else None for i in range(len(headers))}
        try:
            def val(key: str) -> Any:
                idx = mapping.get(key)
                return row[idx] if idx is not None and idx < len(row) else None

            merchant_code = _clean_text(val("merchant_code"))
            product_name = _clean_text(val("product_name"))
            if not merchant_code or merchant_code.upper() == "NA":
                raise ImportValidationError("EMPTY_MERCHANT_CODE", "商家编码不能为空")
            if not product_name or product_name.upper() == "NA":
                raise ImportValidationError("EMPTY_PRODUCT_NAME", "商品名称不能为空")

            previous_name = seen_names.get(merchant_code)
            if previous_name and previous_name != product_name:
                # 商家编码是唯一商品键；商品名称允许随业务调整。
                # 文件内出现不同名称时保留全部有效行，并降级为“名称变体”警告，
                # 后续由管理员在数据质量页聚合到一个规范商品名称。
                parsed.warnings.append(
                    RowError(
                        row_no=row_no,
                        code="PRODUCT_NAME_VARIANT_IN_FILE",
                        message=f"商家编码 {merchant_code} 出现名称变体：{previous_name} / {product_name}；将按商家编码聚合",
                        raw_data=raw,
                    )
                )
            seen_names.setdefault(merchant_code, product_name)

            base: dict[str, Any] = {
                "row_no": row_no,
                "merchant_code": merchant_code,
                "product_name": product_name,
            }

            if data_type == "product":
                base.update({
                    "spec": _clean_text(val("spec")) or None,
                    "brand": _clean_text(val("brand")) or None,
                    "category": _clean_text(val("category")) or None,
                    "barcode": _clean_text(val("barcode")) or None,
                })
            elif data_type == "sales":
                shop = _clean_text(val("shop_name"))
                wh = _clean_text(val("warehouse_name"))
                if not shop or shop.upper() == "NA":
                    raise ImportValidationError("EMPTY_SHOP", "店铺不能为空")
                if not wh or wh.upper() == "NA":
                    raise ImportValidationError("EMPTY_WAREHOUSE", "仓库不能为空")
                base.update({
                    "business_date": business_date,
                    "shop_name": shop,
                    "warehouse_name": wh,
                    "sales_qty": _to_decimal(val("sales_qty"), "销量"),
                    "avg_price": _to_decimal(val("avg_price"), "均价"),
                })
            elif data_type == "inventory":
                wh = _clean_text(val("warehouse_name"))
                if not wh or wh.upper() == "NA":
                    raise ImportValidationError("EMPTY_WAREHOUSE", "仓库不能为空")
                production_raw = val("production_date")
                expire_raw = val("expire_date")
                production = None
                expire = None
                try:
                    production = _to_date(production_raw, "生产日期")
                except ImportValidationError as date_exc:
                    parsed.warnings.append(RowError(row_no=row_no, code="INVALID_PRODUCTION_DATE", message=date_exc.message + "；该批次仍导入库存，效期百分比暂不可计算", raw_data=raw))
                try:
                    expire = _to_date(expire_raw, "过期日期")
                except ImportValidationError as date_exc:
                    parsed.warnings.append(RowError(row_no=row_no, code="INVALID_EXPIRE_DATE", message=date_exc.message + "；该批次仍导入库存，但不参与效期预警", raw_data=raw))
                base.update({
                    "snapshot_date": business_date,
                    "warehouse_name": wh,
                    "stock_qty": _to_decimal(val("stock_qty"), "库存数量"),
                    "production_date": production,
                    "expire_date": expire,
                })
            elif data_type == "aging":
                wh = _clean_text(val("warehouse_name"))
                if not wh or wh.upper() == "NA":
                    raise ImportValidationError("EMPTY_WAREHOUSE", "仓库不能为空")
                base.update({
                    "snapshot_date": business_date,
                    "warehouse_name": wh,
                    "stock_qty": _to_decimal(val("stock_qty"), "库存数量"),
                    "aging_days": _to_int(val("aging_days"), "库龄天数"),
                })
            parsed.rows.append(base)
        except ImportValidationError as exc:
            parsed.errors.append(RowError(row_no=row_no, code=exc.code, message=exc.message, raw_data=raw))

    return parsed


def json_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {k: json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(v) for v in value]
    return value


def preview_commit_policy(parsed: ParsedImport) -> dict[str, bool]:
    """给数据中心前端一个明确的提交策略。

    - 销量/商品资料允许“有效行 + 错误行”的部分导入；
    - 库存效期/库龄是整表最新快照，存在任何错误行时禁止覆盖当前快照；
    - warning 不阻止提交（商品名称变体属于 warning）。
    """
    snapshot_replace_blocked = parsed.data_type in {"inventory", "aging"} and parsed.error_rows > 0
    return {
        "commit_allowed": parsed.valid_rows > 0 and not snapshot_replace_blocked,
        "partial_import": parsed.valid_rows > 0 and parsed.error_rows > 0 and not snapshot_replace_blocked,
    }


def parsed_preview(parsed: ParsedImport, limit: int = 5) -> dict[str, Any]:
    return {
        "summary": parsed.summary(),
        **preview_commit_policy(parsed),
        "snapshot_replace_blocked": parsed.data_type in {"inventory", "aging"} and parsed.error_rows > 0,
        "rows": [json_safe(r) for r in parsed.rows[:limit]],
        "errors": [
            {"row_no": e.row_no, "code": e.code, "message": e.message, "raw_data": json_safe(e.raw_data)}
            for e in parsed.errors[:limit]
        ],
        "warnings": [
            {"row_no": e.row_no, "code": e.code, "message": e.message, "raw_data": json_safe(e.raw_data)}
            for e in parsed.warnings[:limit]
        ],
    }
