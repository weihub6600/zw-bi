from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from .analysis_service import get_expiry_batches, get_inventory_analysis
from .auth_service import record_activity
from .dashboard_service import DashboardScope, _scope_clauses, _where, selected_sales_range
from .permission_service import assert_can_view_department
from .product_category_codec import parse_category_names, serialize_category_names
from ..core.config import settings
from ..core.timezone import today_local


class ExportError(ValueError):
    pass


_DANGEROUS_EXCEL_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_ILLEGAL_XML_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# 各类导出的中文表头（顺序即列顺序）
PRODUCT_HEADERS = ["商家编码", "商品名称", "规格", "品牌", "分类", "条码"]
SALES_HEADERS = ["业务日期", "店铺", "仓库", "商家编码", "商品名称", "销量", "均价", "销售额"]
INVENTORY_HEADERS = ["快照日期", "仓库", "商家编码", "商品名称", "库存数量", "生产日期", "过期日期"]
AGING_HEADERS = ["统计日期", "仓库", "商家编码", "商品名称", "库存数量", "库龄天数"]
# 库存明细（InventoryView 的 SKU 级库存健康度）
INVENTORY_ANALYSIS_HEADERS = ["商家编码", "商品名称", "库存分类", "库存数量", "昨日销量", "近7天销量", "近14天销量", "近30天销量", "上月销量", "预测日均销量", "可售天数", "加权库龄天数", "最大库龄天数"]
# 效期批次明细（ExpiryView 的批次级效期）
EXPIRY_BATCH_HEADERS = ["仓库", "商家编码", "商品名称", "库存数量", "生产日期", "过期日期", "保质总天数", "剩余天数", "剩余效期%", "效期状态", "规则来源"]

# 各类导出对应的文件名前缀（中文，用于 filename* 展示）
FILENAME_PREFIX = {
    "product": "商品资料",
    "sales": "销量明细",
    "inventory": "库存效期",
    "aging": "库龄数据",
    "inventory_analysis": "库存明细",
    "expiry_batches": "效期批次",
}

# 各类导出对应的 ASCII 文件名前缀（用于 Content-Disposition 的 filename 兜底，避免 latin-1 编码失败）
FILENAME_ASCII = {
    "product": "product",
    "sales": "sales",
    "inventory": "inventory_expiry",
    "aging": "aging",
    "inventory_analysis": "inventory",
    "expiry_batches": "expiry",
}


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    return float(value)


def _safe_excel_cell(value: Any) -> tuple[Any, bool]:
    """Return a cleaned value and whether Excel must treat it as text.

    Numeric/date values retain their native types. External text is cleaned of
    XML 1.0-illegal controls, and formula-like prefixes are explicitly marked
    for string serialization by the row writer without deleting user content.
    """
    if not isinstance(value, str):
        return value, False
    cleaned = _ILLEGAL_XML_CHARS_RE.sub("", value)
    return cleaned, cleaned.startswith(_DANGEROUS_EXCEL_PREFIXES)


def _append_safe_row(ws, row: Iterable[Any], *, write_only: bool) -> None:
    if write_only:
        cells = []
        for value in row:
            cleaned, force_text = _safe_excel_cell(value)
            cell = WriteOnlyCell(ws, value=cleaned)
            if force_text:
                cell.data_type = "s"
                cell.quotePrefix = True
            cells.append(cell)
        ws.append(cells)
        return

    prepared = [_safe_excel_cell(value) for value in row]
    ws.append([value for value, _force_text in prepared])
    row_index = ws.max_row
    for column_index, (_value, force_text) in enumerate(prepared, start=1):
        cell = ws.cell(row=row_index, column=column_index)
        if force_text:
            cell.data_type = "s"
            cell.quotePrefix = True


def _ensure_export_row_limit(row_count: int) -> None:
    if row_count > settings.max_export_rows:
        raise ExportError(
            f"当前结果超过最大导出行数 {settings.max_export_rows}，请缩小日期范围或增加筛选条件"
        )


def _make_workbook(headers: list[str], rows: list[list[Any]], date_cols: set[int], *, write_only: bool = False) -> bytes:
    _ensure_export_row_limit(len(rows))
    if write_only:
        # 流式写入，适用于大数据量（如效期批次完整导出），降低内存占用。
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("导出数据")
        _append_safe_row(ws, headers, write_only=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        for col_idx, _ in enumerate(headers, start=1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = 12 if col_idx in date_cols else 18
        for row in rows:
            _append_safe_row(ws, row, write_only=True)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    wb = Workbook()
    ws = wb.active
    ws.title = "导出数据"
    _append_safe_row(ws, headers, write_only=False)
    # 冻结首行、自动筛选
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    # 表头加粗
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        _append_safe_row(ws, row, write_only=False)
    # 列宽：按内容自适应（封顶 40），日期列固定 12
    for col_idx, _ in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        if col_idx in date_cols:
            ws.column_dimensions[letter].width = 12
        else:
            max_len = len(str(headers[col_idx - 1]))
            for r in range(2, len(rows) + 2):
                cell = ws.cell(row=r, column=col_idx)
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max_len + 2, 40)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _product_rows(db: Session, department_id: int) -> list[list[Any]]:
    """仅导出与当前部门相关的商品，严禁 SELECT * FROM products 造成跨部门泄漏。"""
    rows = db.execute(
        text(
            """
            SELECT p.id, p.merchant_code, p.product_name, p.spec, p.brand, p.category, p.barcode
            FROM products p
            WHERE EXISTS (
                SELECT 1 FROM sales_daily s WHERE s.department_id=:department_id AND s.product_id=p.id
            )
               OR EXISTS (
                SELECT 1 FROM inventory_batch i WHERE i.department_id=:department_id AND i.product_id=p.id
            )
               OR EXISTS (
                SELECT 1 FROM aging_snapshot a WHERE a.department_id=:department_id AND a.product_id=p.id
            )
               OR EXISTS (
                SELECT 1 FROM import_batches b
                WHERE b.id=p.import_batch_id AND b.department_id=:department_id
            )
            ORDER BY p.merchant_code
            LIMIT :row_limit
            """
        ),
        {"department_id": department_id, "row_limit": settings.max_export_rows + 1},
    ).all()
    product_ids = [int(row[0]) for row in rows]
    relation_names: dict[int, list[str]] = {}
    known_names: list[str] = []
    try:
        known_names = [str(value) for value in db.execute(text("SELECT category_name FROM product_categories")).scalars().all()]
        if product_ids:
            placeholders = ",".join(f":product_{i}" for i in range(len(product_ids)))
            relation_rows = db.execute(
                text(
                    f"SELECT pcr.product_id, pc.category_name "
                    f"FROM product_category_relations pcr JOIN product_categories pc ON pc.id=pcr.category_id "
                    f"WHERE pcr.product_id IN ({placeholders}) ORDER BY pcr.product_id, pc.category_name"
                ),
                {f"product_{i}": value for i, value in enumerate(product_ids)},
            ).all()
            for product_id, category_name in relation_rows:
                relation_names.setdefault(int(product_id), []).append(str(category_name))
    except Exception:
        # Older restored databases may not have the relation tables yet.
        relation_names = {}

    result: list[list[Any]] = []
    for row in rows:
        product_id, merchant_code, product_name, spec, brand, legacy_category, barcode = row
        names = relation_names.get(int(product_id))
        if names is None:
            names = parse_category_names(legacy_category, known_names=known_names)
        result.append([merchant_code, product_name, spec, brand, serialize_category_names(names), barcode])
    return result


def _get_latest_sales_date(db: Session, department_id: int) -> date | None:
    value = db.execute(
        text("SELECT MAX(business_date) FROM sales_daily WHERE department_id=:department_id"),
        {"department_id": department_id},
    ).scalar_one_or_none()
    if value is None:
        return None
    if hasattr(value, "year"):
        return value
    # SQLite 下 business_date 以字符串返回，统一为 date。
    return date.fromisoformat(str(value)[:10])


def _sales_rows(db: Session, department_id: int, scope: DashboardScope) -> list[list[Any]]:
    """销量明细导出，复用 dashboard/analysis 同一套筛选逻辑（_scope_clauses）。

    返回完整筛选结果（不分页），页面列表与该结果唯一差异仅在于分页。
    """
    latest = _get_latest_sales_date(db, department_id)
    if latest is None:
        return []
    # 复用 selected_sales_range 解析日期区间（自定义区间或预设锚点）。
    start_date, end_date, _days = selected_sales_range(scope, latest)

    params: dict[str, Any] = {
        "department_id": department_id,
        "start_date": start_date,
        "end_date": end_date,
    }
    filters = _scope_clauses(scope, params, product_alias="p", shop_alias="s", warehouse_alias="w")
    rows = db.execute(
        text(
            f"""
            SELECT sd.business_date, s.source_name AS shop, w.source_name AS warehouse,
                   p.merchant_code, p.product_name, sd.sales_qty, sd.avg_price,
                   sd.sales_qty * sd.avg_price AS amount
            FROM sales_daily sd
            JOIN products p ON p.id=sd.product_id
            JOIN shops s ON s.id=sd.shop_id
            JOIN warehouses w ON w.id=sd.warehouse_id
            WHERE sd.department_id=:department_id
              AND sd.business_date BETWEEN :start_date AND :end_date
              {_where(filters)}
            ORDER BY sd.business_date, s.source_name, w.source_name, p.merchant_code
            LIMIT :row_limit
            """
        ),
        {**params, "row_limit": settings.max_export_rows + 1},
    ).all()
    out: list[list[Any]] = []
    for r in rows:
        business_date, shop, wh, code, name, qty, price, amount = r
        out.append([
            business_date if hasattr(business_date, "year") else business_date,
            shop, wh, code, name, _to_float(qty), _to_float(price), _to_float(amount),
        ])
    return out


def _latest_snapshot_date(db: Session, department_id: int, table: str) -> date | None:
    value = db.execute(
        text(f"SELECT MAX(snapshot_date) FROM {table} WHERE department_id=:department_id"),
        {"department_id": department_id},
    ).scalar()
    return value if value is not None else None


def _inventory_rows(db: Session, department_id: int) -> list[list[Any]]:
    latest = _latest_snapshot_date(db, department_id, "inventory_batch")
    if latest is None:
        return []
    rows = db.execute(
        text(
            """
            SELECT ib.snapshot_date, w.source_name, p.merchant_code, p.product_name,
                   ib.stock_qty, ib.production_date, ib.expire_date
            FROM inventory_batch ib
            JOIN products p ON p.id=ib.product_id
            JOIN warehouses w ON w.id=ib.warehouse_id
            WHERE ib.department_id=:department_id AND ib.snapshot_date=:snapshot_date
            ORDER BY w.source_name, p.merchant_code
            LIMIT :row_limit
            """
        ),
        {
            "department_id": department_id,
            "snapshot_date": latest,
            "row_limit": settings.max_export_rows + 1,
        },
    ).all()
    out: list[list[Any]] = []
    for r in rows:
        snapshot, wh, code, name, qty, prod, expire = r
        out.append([
            snapshot if hasattr(snapshot, "year") else snapshot,
            wh, code, name, _to_float(qty),
            prod if hasattr(prod, "year") else prod,
            expire if hasattr(expire, "year") else expire,
        ])
    return out


def _aging_rows(db: Session, department_id: int) -> list[list[Any]]:
    latest = _latest_snapshot_date(db, department_id, "aging_snapshot")
    if latest is None:
        return []
    rows = db.execute(
        text(
            """
            SELECT ag.snapshot_date, w.source_name, p.merchant_code, p.product_name,
                   ag.stock_qty, ag.aging_days
            FROM aging_snapshot ag
            JOIN products p ON p.id=ag.product_id
            JOIN warehouses w ON w.id=ag.warehouse_id
            WHERE ag.department_id=:department_id AND ag.snapshot_date=:snapshot_date
            ORDER BY w.source_name, p.merchant_code
            LIMIT :row_limit
            """
        ),
        {
            "department_id": department_id,
            "snapshot_date": latest,
            "row_limit": settings.max_export_rows + 1,
        },
    ).all()
    out: list[list[Any]] = []
    for r in rows:
        snapshot, wh, code, name, qty, days = r
        out.append([
            snapshot if hasattr(snapshot, "year") else snapshot,
            wh, code, name, _to_float(qty), int(days),
        ])
    return out


def _build_xlsx(
    data_type: str,
    rows: list[list[Any]],
    department_code: str,
    *,
    write_only: bool = False,
    headers_override: list[str] | None = None,
) -> tuple[bytes, str, int, str]:
    _ensure_export_row_limit(len(rows))
    if not rows:
        # 空结果：不生成无意义 Excel，也不写 data_export 审计记录。
        raise ExportError("当前筛选条件下暂无可导出数据")
    if data_type == "product":
        headers, date_cols = PRODUCT_HEADERS, set()
    elif data_type == "sales":
        headers, date_cols = SALES_HEADERS, {1}
    elif data_type == "inventory":
        headers, date_cols = INVENTORY_HEADERS, {1, 6, 7}
    elif data_type == "aging":
        headers, date_cols = AGING_HEADERS, {1}
    elif data_type == "inventory_analysis":
        headers, date_cols = INVENTORY_ANALYSIS_HEADERS, set()
    elif data_type == "expiry_batches":
        headers, date_cols = EXPIRY_BATCH_HEADERS, {5, 6}
    else:
        raise ExportError(f"不支持的数据类型：{data_type}")
    content = _make_workbook(headers_override or headers, rows, date_cols, write_only=write_only)
    suffix = today_local().strftime('%Y%m%d')
    filename = f"{department_code}_{FILENAME_PREFIX[data_type]}_{suffix}.xlsx"
    ascii_filename = f"{department_code}_{FILENAME_ASCII[data_type]}_{suffix}.xlsx"
    return content, filename, len(rows), ascii_filename


def export_product(db: Session, actor_user_id: str, department_code: str) -> dict[str, Any]:
    actor, department = assert_can_view_department(db, actor_user_id, department_code)
    rows = _product_rows(db, int(department["id"]))
    content, filename, row_count, ascii_filename = _build_xlsx("product", rows, department["code"])
    record_activity(
        db, int(actor["id"]), "data_export",
        department_id=int(department["id"]),
        detail={"department_code": department["code"], "data_type": "product", "start_date": None, "end_date": None, "row_count": row_count, "filename": filename},
    )
    return {"content": content, "filename": filename, "ascii_filename": ascii_filename, "row_count": row_count}


def export_sales(
    db: Session,
    actor_user_id: str,
    department_code: str,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    days: int = 30,
    shops: Iterable[str] = (),
    warehouses: Iterable[str] = (),
    product_search: str = "",
    include_name: str = "",
    exclude_name: str = "",
    product_codes: Iterable[str] = (),
) -> dict[str, Any]:
    """导出当前筛选结果对应的全部销量明细（不分页）。

    筛选条件与 dashboard/analysis 列表完全一致（复用 _scope_clauses），
    唯一差异是导出返回完整结果而不分页。
    日期：自定义区间优先；否则按 days 预设（1/7/14/30）以最新销售日为锚点。
    """
    actor, department = assert_can_view_department(db, actor_user_id, department_code)
    if start_date and end_date and start_date > end_date:
        raise ExportError("开始日期不能晚于结束日期")
    scope = DashboardScope(
        actor_user_id=actor_user_id,
        department_code=department_code,
        days=days,
        start_date=start_date,
        end_date=end_date,
        shops=tuple(shops or ()),
        warehouses=tuple(warehouses or ()),
        product_search=product_search or "",
        include_name=include_name or "",
        exclude_name=exclude_name or "",
        product_codes=tuple(product_codes or ()),
    )
    rows = _sales_rows(db, int(department["id"]), scope)
    content, filename, row_count, ascii_filename = _build_xlsx("sales", rows, department["code"])
    record_activity(
        db, int(actor["id"]), "data_export",
        department_id=int(department["id"]),
        detail={"department_code": department["code"], "data_type": "sales", "start_date": start_date.isoformat() if start_date else None, "end_date": end_date.isoformat() if end_date else None, "row_count": row_count, "filename": filename},
    )
    return {"content": content, "filename": filename, "ascii_filename": ascii_filename, "row_count": row_count}


def _scope_from_params(actor_user_id: str, department_code: str, *, warehouses, product_search, include_name, exclude_name, product_codes, product_category_ids=(), start_date=None, end_date=None, shops=(), days=30) -> DashboardScope:
    return DashboardScope(
        actor_user_id=actor_user_id,
        department_code=department_code,
        days=days,
        start_date=start_date,
        end_date=end_date,
        shops=tuple(shops or ()),
        warehouses=tuple(warehouses or ()),
        product_search=product_search or "",
        include_name=include_name or "",
        exclude_name=exclude_name or "",
        product_codes=tuple(product_codes or ()),
        product_category_ids=tuple(product_category_ids or ()),
    )


def export_inventory_analysis(
    db: Session,
    actor_user_id: str,
    department_code: str,
    *,
    shops: Iterable[str] = (),
    warehouses: Iterable[str] = (),
    days: int = 30,
    product_search: str = "",
    include_name: str = "",
    exclude_name: str = "",
    product_codes: Iterable[str] = (),
    product_category_ids: Iterable[int] = (),
    category: str | None = None,
    detail_warehouses: bool = False,
    detail_product_categories: bool = False,
) -> dict[str, Any]:
    """导出库存明细（InventoryView 的 SKU 级库存健康度），复用 get_inventory_analysis 完整筛选与权限。"""
    scope = _scope_from_params(
        actor_user_id, department_code,
        shops=shops, warehouses=warehouses, days=days,
        product_search=product_search, include_name=include_name,
        exclude_name=exclude_name, product_codes=product_codes,
        product_category_ids=product_category_ids,
    )
    result = get_inventory_analysis(
        db, scope, category=category, detail_warehouses=detail_warehouses,
        detail_product_categories=detail_product_categories,
    )
    actor, department = assert_can_view_department(db, actor_user_id, department_code)
    rows = result["rows"]
    _ensure_export_row_limit(len(rows))
    warehouse_columns = result.get("meta", {}).get("warehouse_columns", [])
    headers = INVENTORY_ANALYSIS_HEADERS
    if detail_product_categories:
        headers = ["商家编码", "商品名称", "商品分类", "库存分类", "库存数量", "昨日销量", "近7天销量", "近14天销量", "近30天销量", "上月销量", "预测日均销量", "可售天数", "加权库龄天数", "最大库龄天数"]
    if detail_warehouses:
        headers = [
            "商家编码", "商品名称",
            *( ["商品分类"] if detail_product_categories else [] ),
            "库存分类", *warehouse_columns,
            "总库存", "昨日销量", "近7天销量", "近14天销量", "近30天销量", "上月销量", "预测日均销量",
            "可售天数", "加权库龄天数", "最大库龄天数",
        ]
    data = [
        [
            r["sku"], r["name"],
            *([r.get("product_category_names", "")] if detail_product_categories else []),
            r["category_label"], _to_float(r["stock_qty"]),
            _to_float(r.get("sales_yesterday")), _to_float(r["sales7"]), _to_float(r["sales14"]), _to_float(r["sales30"]), _to_float(r.get("sales_last_month")),
            _to_float(r["predicted_daily"]), _to_float(r["cover_days"]),
            _to_float(r["weighted_aging_days"]), _to_float(r["max_aging_days"]),
        ]
        for r in rows
    ]
    if detail_warehouses:
        data = [
            [
                r["sku"], r["name"],
                *([r.get("product_category_names", "")] if detail_product_categories else []),
                r["category_label"],
                *[_to_float(r.get("warehouse_stocks", {}).get(warehouse, 0)) for warehouse in warehouse_columns],
                _to_float(r["stock_qty"]), _to_float(r.get("sales_yesterday")), _to_float(r["sales7"]), _to_float(r["sales14"]),
                _to_float(r["sales30"]), _to_float(r.get("sales_last_month")), _to_float(r["predicted_daily"]),
                _to_float(r["cover_days"]), _to_float(r["weighted_aging_days"]),
                _to_float(r["max_aging_days"]),
            ]
            for r in rows
        ]
    content, filename, row_count, ascii_filename = _build_xlsx(
        "inventory_analysis", data, department["code"], headers_override=headers
    )
    record_activity(
        db, int(actor["id"]), "data_export",
        department_id=int(department["id"]),
        detail={"department_code": department["code"], "data_type": "inventory_analysis", "start_date": None, "end_date": None, "row_count": row_count, "filename": filename},
    )
    return {"content": content, "filename": filename, "ascii_filename": ascii_filename, "row_count": row_count}


def export_expiry_batches(
    db: Session,
    actor_user_id: str,
    department_code: str,
    *,
    warehouses: Iterable[str] = (),
    product_search: str = "",
    include_name: str = "",
    exclude_name: str = "",
    product_codes: Iterable[str] = (),
    product_category_ids: Iterable[int] = (),
    statuses: Iterable[str] = (),
    remaining_days_min: int | None = None,
    remaining_days_max: int | None = None,
    remaining_pct_min: float | None = None,
    remaining_pct_max: float | None = None,
    merchant_code: str | None = None,
) -> dict[str, Any]:
    """导出效期批次明细（ExpiryView），复用 get_expiry_batches 完整筛选与权限。"""
    scope = _scope_from_params(
        actor_user_id, department_code,
        warehouses=warehouses, product_search=product_search, include_name=include_name,
        exclude_name=exclude_name, product_codes=product_codes,
        product_category_ids=product_category_ids,
    )
    result = get_expiry_batches(
        db, scope,
        statuses=tuple(statuses or ()),
        remaining_days_min=remaining_days_min,
        remaining_days_max=remaining_days_max,
        remaining_pct_min=remaining_pct_min,
        remaining_pct_max=remaining_pct_max,
        merchant_code=merchant_code,
        limit=None,  # 完整导出，不截断
    )
    actor, department = assert_can_view_department(db, actor_user_id, department_code)
    _ensure_export_row_limit(int(result.get("filtered_count", len(result["rows"]))))
    rows = result["rows"]
    data = [
        [
            r["warehouse"], r["sku"], r["name"], _to_float(r["stock_qty"]),
            r["production_date"], r["expire_date"], r["total_shelf_days"],
            r["remaining_days"], _to_float(r["remaining_pct"]), r["status"], r["rule_source"],
        ]
        for r in rows
    ]
    content, filename, row_count, ascii_filename = _build_xlsx("expiry_batches", data, department["code"], write_only=True)
    record_activity(
        db, int(actor["id"]), "data_export",
        department_id=int(department["id"]),
        detail={"department_code": department["code"], "data_type": "expiry_batches", "start_date": None, "end_date": None, "row_count": row_count, "filename": filename},
    )
    return {"content": content, "filename": filename, "ascii_filename": ascii_filename, "row_count": row_count}


def export_inventory(db: Session, actor_user_id: str, department_code: str) -> dict[str, Any]:
    actor, department = assert_can_view_department(db, actor_user_id, department_code)
    rows = _inventory_rows(db, int(department["id"]))
    content, filename, row_count, ascii_filename = _build_xlsx("inventory", rows, department["code"])
    record_activity(
        db, int(actor["id"]), "data_export",
        department_id=int(department["id"]),
        detail={"department_code": department["code"], "data_type": "inventory", "start_date": None, "end_date": None, "row_count": row_count, "filename": filename},
    )
    return {"content": content, "filename": filename, "ascii_filename": ascii_filename, "row_count": row_count}


def export_aging(db: Session, actor_user_id: str, department_code: str) -> dict[str, Any]:
    actor, department = assert_can_view_department(db, actor_user_id, department_code)
    rows = _aging_rows(db, int(department["id"]))
    content, filename, row_count, ascii_filename = _build_xlsx("aging", rows, department["code"])
    record_activity(
        db, int(actor["id"]), "data_export",
        department_id=int(department["id"]),
        detail={"department_code": department["code"], "data_type": "aging", "start_date": None, "end_date": None, "row_count": row_count, "filename": filename},
    )
    return {"content": content, "filename": filename, "ascii_filename": ascii_filename, "row_count": row_count}
