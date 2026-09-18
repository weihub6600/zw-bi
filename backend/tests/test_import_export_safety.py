from __future__ import annotations

import os
import tempfile
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest import mock

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook

from app.routes import imports as import_routes
from app.services import export_service, import_parser
from app.services.export_service import ExportError, _make_workbook
from app.services.import_parser import ImportValidationError, _clean_text, _to_decimal, parse_import_file


class ExcelExportSafetyTests(unittest.TestCase):
    def _load_values(self, content: bytes):
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.addCleanup(workbook.close)
        return workbook.active

    def test_formula_like_text_is_not_serialized_as_formula(self):
        dangerous = [
            "=1+1",
            "+1+1",
            "-1+1",
            "@SUM(1,2)",
            '=HYPERLINK("https://example.com","click")',
            "\t=1+1",
            "\r=1+1",
        ]
        for write_only in (False, True):
            with self.subTest(write_only=write_only):
                sheet = self._load_values(
                    _make_workbook(
                        ["文本", "整数", "小数"],
                        [[value, 123, 12.5] for value in dangerous],
                        set(),
                        write_only=write_only,
                    )
                )
                for row_index, expected in enumerate(dangerous, start=2):
                    cell = sheet.cell(row=row_index, column=1)
                    # XML parsers normalize carriage returns to line feeds;
                    # the content remains text and is not dropped/executed.
                    self.assertEqual(cell.value, expected.replace("\r", "\n"))
                    self.assertNotEqual(cell.data_type, "f")
                    self.assertTrue(cell.quotePrefix)
                self.assertEqual(sheet["B2"].data_type, "n")
                self.assertEqual(sheet["C2"].data_type, "n")

    def test_illegal_xml_controls_are_removed_but_legal_whitespace_is_preserved(self):
        text = "中文\x00A\x07B\x0bC\n第二行\t末尾\r回车"
        for write_only in (False, True):
            with self.subTest(write_only=write_only):
                sheet = self._load_values(
                    _make_workbook(["文本"], [[text]], set(), write_only=write_only)
                )
                self.assertEqual(sheet["A2"].value, "中文ABC\n第二行\t末尾\n回车")

    def test_export_limit_rejects_before_workbook_creation(self):
        with mock.patch.object(export_service.settings, "max_export_rows", 1), mock.patch.object(
            export_service, "Workbook"
        ) as workbook:
            with self.assertRaises(ExportError) as ctx:
                _make_workbook(["文本"], [["A"], ["B"]], set())
        workbook.assert_not_called()
        self.assertIn("缩小日期范围", str(ctx.exception))


class ImportParserSafetyTests(unittest.TestCase):
    def _temp_path(self, suffix: str) -> Path:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        return Path(directory.name) / f"input{suffix}"

    def _write_xlsx(self, rows: list[list[object]]) -> Path:
        path = self._temp_path(".xlsx")
        workbook = Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_csv_quoted_multiline_field_is_one_business_row(self):
        path = self._temp_path(".csv")
        path.write_text(
            '商家编码,商品名称\nA001,"商品名称第一行\n商品名称第二行"\n',
            encoding="utf-8",
            newline="",
        )
        parsed = parse_import_file(path, "product")
        self.assertEqual(parsed.source_row_count, 1)
        self.assertEqual(parsed.valid_rows, 1)
        self.assertEqual(parsed.rows[0]["product_name"], "商品名称第一行\n商品名称第二行")

    def test_utf8_bom_does_not_pollute_first_header(self):
        path = self._temp_path(".csv")
        path.write_text("商家编码,商品名称\nA001,商品A\n", encoding="utf-8-sig")
        parsed = parse_import_file(path, "product")
        self.assertEqual(parsed.valid_rows, 1)
        self.assertIn("merchant_code", parsed.header_mapping)

    def test_csv_row_limit_stops_with_business_error(self):
        path = self._temp_path(".csv")
        path.write_text("商家编码,商品名称\nA001,商品A\nA002,商品B\n", encoding="utf-8")
        with mock.patch.object(import_parser.settings, "max_import_rows", 1):
            with self.assertRaises(ImportValidationError) as ctx:
                parse_import_file(path, "product")
        self.assertEqual(ctx.exception.code, "TOO_MANY_ROWS")
        self.assertIn("1 行上限", ctx.exception.message)

    def test_xlsx_row_limit_stops_with_business_error(self):
        path = self._write_xlsx(
            [["商家编码", "商品名称"], ["A001", "商品A"], ["A002", "商品B"]]
        )
        with mock.patch.object(import_parser.settings, "max_import_rows", 1):
            with self.assertRaises(ImportValidationError) as ctx:
                parse_import_file(path, "product")
        self.assertEqual(ctx.exception.code, "TOO_MANY_ROWS")

    def test_xlsx_worksheet_limit_is_enforced(self):
        path = self._temp_path(".xlsx")
        workbook = Workbook()
        workbook.active.append(["商家编码", "商品名称"])
        workbook.create_sheet("extra")
        workbook.save(path)
        workbook.close()
        with mock.patch.object(import_parser.settings, "max_import_worksheets", 1):
            with self.assertRaises(ImportValidationError) as ctx:
                parse_import_file(path, "product")
        self.assertEqual(ctx.exception.code, "TOO_MANY_WORKSHEETS")

    def test_xlsx_expanded_size_limit_is_enforced(self):
        path = self._write_xlsx([["商家编码", "商品名称"], ["A001", "商品A"]])
        with mock.patch.object(import_parser.settings, "max_xlsx_uncompressed_bytes", 1):
            with self.assertRaises(ImportValidationError) as ctx:
                parse_import_file(path, "product")
        self.assertEqual(ctx.exception.code, "XLSX_EXPANDED_TOO_LARGE")

    def test_import_text_removes_only_illegal_xml_controls(self):
        self.assertEqual(_clean_text("中文\x00A\x01B\x07C\x0bD\x0cE\x1fF\n第二行\t末尾"), "中文ABCDEF\n第二行\t末尾")

    def test_non_finite_and_excessive_decimals_are_rejected(self):
        for value in ("NaN", "Infinity", "-inf", "1E+1000"):
            with self.subTest(value=value), self.assertRaises(ImportValidationError):
                _to_decimal(value, "销量")


class UploadLimitTests(unittest.TestCase):
    def test_upload_below_limit_is_saved(self):
        upload = UploadFile(filename="small.csv", file=BytesIO(b"abc"))
        with mock.patch.object(import_routes.settings, "max_upload_bytes", 4):
            path = import_routes._save_upload(upload)
        try:
            self.assertEqual(Path(path).read_bytes(), b"abc")
        finally:
            os.unlink(path)

    def test_oversized_upload_returns_413_deletes_partial_and_skips_import(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        captured_path = Path(directory.name) / "oversized.csv"

        def make_temp(**_kwargs):
            return os.open(captured_path, os.O_CREAT | os.O_RDWR), str(captured_path)

        upload = UploadFile(filename="large.csv", file=BytesIO(b"12345"))
        with mock.patch.object(import_routes.settings, "max_upload_bytes", 4), mock.patch.object(
            import_routes.tempfile, "mkstemp", side_effect=make_temp
        ), mock.patch.object(import_routes, "import_file") as import_file:
            with self.assertRaises(HTTPException) as ctx:
                import_routes.commit_import(
                    data_type="product",
                    department_code="B2C",
                    business_date=None,
                    file=upload,
                    actor={"user_id": "ADMIN"},
                    db=mock.MagicMock(),
                )
        self.assertEqual(ctx.exception.status_code, 413)
        self.assertFalse(captured_path.exists())
        import_file.assert_not_called()


if __name__ == "__main__":
    unittest.main()
