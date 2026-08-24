# -*- coding: utf-8 -*-
from __future__ import annotations

import io
import unittest
from urllib.parse import unquote

from openpyxl import load_workbook

from app.routes.exports import _xlsx_response
from app.services.export_service import _make_workbook


def _sample_xlsx():
    return _make_workbook(["\u5546\u5bb6\u7f16\u7801"], [["A001"], ["A002"]], set())  # 商家编码


class ContentDispositionTests(unittest.TestCase):
    def test_chinese_filename_no_unicode_error_and_rfc5987(self):
        body = _sample_xlsx()
        filename = "B2C_\u9500\u91cf\u660e\u7ec6_20260824.xlsx"  # B2C_销量明细_...
        ascii_name = "B2C_sales_20260824.xlsx"
        resp = _xlsx_response(body, filename, row_count=2, ascii_filename=ascii_name)

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        cd = resp.headers["Content-Disposition"]
        # ASCII fallback 在 filename=，中文在 filename*
        self.assertIn('filename="B2C_sales_20260824.xlsx"', cd)
        self.assertIn("filename*=UTF-8''", cd)
        # filename* 解出的中文名正确
        star = cd.split("filename*=UTF-8''")[1]
        decoded = unquote(star)
        self.assertEqual(decoded, filename)
        # body 可被 openpyxl 打开
        wb = load_workbook(io.BytesIO(resp.body))
        self.assertEqual(wb.active.max_row, 3)  # 表头 + 2 行

    def test_ascii_fallback_default_when_missing(self):
        body = _sample_xlsx()
        resp = _xlsx_response(body, "B2C_\u5e93\u5b58\u660e\u7ec6_20260824.xlsx")  # 库存明细
        cd = resp.headers["Content-Disposition"]
        self.assertIn('filename="export.xlsx"', cd)
        self.assertIn("filename*=UTF-8''", cd)

    def test_row_count_header(self):
        body = _sample_xlsx()
        resp = _xlsx_response(body, "x.xlsx", row_count=7, ascii_filename="x.xlsx")
        self.assertEqual(resp.headers.get("X-Export-Row-Count"), "7")


if __name__ == "__main__":
    unittest.main()
